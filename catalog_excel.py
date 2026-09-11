from __future__ import annotations

import hashlib
import os
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Mapping

from catalog_store import (
    CATALOG_FIELDS,
    CatalogProduct,
    CatalogStore,
    CatalogValidationError,
    ImportResult,
    normalize_gtin,
)


PRIMARY_CATALOG_FILENAME = "product_catalog.xlsx"
LEGACY_CATALOG_FILENAME = "Список товаров.xlsx"

FIELD_TO_HEADER = {
    "gtin": "GTIN",
    "name": "Наименование",
    "brand": "Бренд",
    "weight": "Вес",
    "shelf_life": "Срок годности",
    "permit_document": "Разрешительная документация",
    "permit_from": "от",
    "permit_to": "до",
    "packaging": "Упаковка",
    "material": "Материал",
    "marking_method": "Нанесение кодов",
    "tnved": "ТН ВЭД",
    "short_name": "ShortName",
    "print_template": "Шаблон печати",
    "client": "Клиент",
    "display_packaging": "Упаковка отображаемая",
}
CATALOG_HEADERS = tuple(FIELD_TO_HEADER[field] for field in CATALOG_FIELDS)


class CatalogImportValidationError(ValueError):
    pass


@dataclass(frozen=True)
class ImportIssue:
    code: str
    message: str
    row_number: int | None = None
    header: str = ""


@dataclass(frozen=True)
class DuplicateGtin:
    gtin: str
    row_numbers: tuple[int, ...]


@dataclass(frozen=True)
class PreparedExcelImport:
    source_path: Path
    source_sha256: str
    records: tuple[dict[str, str], ...]
    present_fields: frozenset[str]
    has_short_col: bool
    blank_rows: int
    duplicates: tuple[DuplicateGtin, ...]
    unknown_headers: tuple[str, ...]
    errors: tuple[ImportIssue, ...]

    @property
    def valid(self) -> bool:
        return not self.errors

    def require_valid(self) -> tuple[dict[str, str], ...]:
        if self.errors:
            details = "; ".join(issue.message for issue in self.errors[:5])
            raise CatalogImportValidationError(details or "Excel catalog is invalid")
        return self.records


def _normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower().replace("ё", "е")


_HEADER_ALIASES = {
    "gtin": ("GTIN", "ГТИН"),
    "name": ("Наименование", "Name"),
    "brand": ("Бренд", "Brand"),
    "weight": ("Вес", "Weight", "Объем", "Объём"),
    "shelf_life": ("Срок годности",),
    "permit_document": (
        "Разрешительная документация",
        "Разрешительная док",
    ),
    "permit_from": ("от",),
    "permit_to": ("до",),
    "packaging": ("Упаковка",),
    "material": ("Материал",),
    "marking_method": ("Нанесение кодов",),
    "tnved": ("ТН ВЭД", "ТНВЭД"),
    "short_name": (
        "ShortName",
        "Short Name",
        "Короткое имя",
    ),
    "print_template": (
        "Шаблон печати",
        "Шаблон",
        "Template",
    ),
    "client": ("Клиент", "Client"),
    "display_packaging": (
        "Упаковка отображаемая",
        "Упаковка для стикера",
        "PACKAGING",
    ),
}

_ALIAS_TO_FIELD = {
    _normalize_header(alias): field
    for field, aliases in _HEADER_ALIASES.items()
    for alias in aliases
}


def _resolve_columns(headers: list[str]):
    columns: dict[str, list[int]] = {}
    unknown: list[str] = []
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if not normalized:
            continue
        field = _ALIAS_TO_FIELD.get(normalized)
        if field is None:
            unknown.append(header)
            continue
        columns.setdefault(field, []).append(index)

    for index, header in enumerate(headers):
        if header not in unknown:
            continue
        normalized = _normalize_header(header)
        fallback_field = None
        if any(token in normalized for token in ("вес", "weight", "объем")):
            fallback_field = "weight"
        elif "бренд" in normalized or "brand" in normalized:
            fallback_field = "brand"
        if fallback_field:
            columns.setdefault(fallback_field, []).append(index)
            unknown.remove(header)
    return columns, unknown


def _cell_text(cell) -> str:
    if cell.data_type == "f":
        raise CatalogImportValidationError("formulas are not supported")
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _gtin_cell_value(cell):
    if cell.data_type == "f":
        raise CatalogValidationError("GTIN formulas are not supported")
    value = cell.value
    if isinstance(value, bool):
        raise CatalogValidationError("GTIN must contain digits")
    if isinstance(value, float) and not value.is_integer():
        raise CatalogValidationError("GTIN must not be fractional")
    return value


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def prepare_excel_import(path: str | os.PathLike[str]) -> PreparedExcelImport:
    source_path = Path(path)
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel catalog import") from exc

    workbook = load_workbook(source_path, data_only=False, read_only=True)
    try:
        worksheet = workbook.active
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1), ())
        headers = ["" if cell.value is None else str(cell.value).strip() for cell in first_row]
        columns, unknown_headers = _resolve_columns(headers)
        errors: list[ImportIssue] = []
        for field, indexes in columns.items():
            if len(indexes) > 1:
                names = ", ".join(headers[index] for index in indexes)
                errors.append(
                    ImportIssue(
                        code="ambiguous_header",
                        message=f"Several columns map to {FIELD_TO_HEADER[field]}: {names}",
                        header=FIELD_TO_HEADER[field],
                    )
                )
        if "gtin" not in columns:
            errors.append(
                ImportIssue(
                    code="missing_gtin_header",
                    message="Required GTIN column is missing",
                    header="GTIN",
                )
            )

        present_fields = frozenset(columns)
        if errors:
            return PreparedExcelImport(
                source_path=source_path,
                source_sha256=_file_sha256(source_path),
                records=(),
                present_fields=present_fields,
                has_short_col="short_name" in columns,
                blank_rows=0,
                duplicates=(),
                unknown_headers=tuple(unknown_headers),
                errors=tuple(errors),
            )

        records_by_gtin: dict[str, dict[str, str]] = {}
        row_numbers_by_gtin: dict[str, list[int]] = {}
        blank_rows = 0
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
            if all(
                cell.value is None or str(cell.value).strip() == ""
                for cell in row
            ):
                blank_rows += 1
                continue
            try:
                gtin_index = columns["gtin"][0]
                gtin_cell = row[gtin_index] if gtin_index < len(row) else None
                if gtin_cell is None:
                    raise CatalogValidationError("GTIN is required")
                gtin = normalize_gtin(_gtin_cell_value(gtin_cell))
                values = {field: "" for field in CATALOG_FIELDS}
                values["gtin"] = gtin
                for field, indexes in columns.items():
                    if field == "gtin":
                        continue
                    index = indexes[0]
                    values[field] = _cell_text(row[index]) if index < len(row) else ""
            except (CatalogValidationError, CatalogImportValidationError) as exc:
                errors.append(
                    ImportIssue(
                        code="invalid_row",
                        message=f"Row {row_number}: {exc}",
                        row_number=row_number,
                    )
                )
                continue
            row_numbers_by_gtin.setdefault(gtin, []).append(row_number)
            records_by_gtin[gtin] = values

        duplicates = tuple(
            DuplicateGtin(gtin=gtin, row_numbers=tuple(row_numbers))
            for gtin, row_numbers in row_numbers_by_gtin.items()
            if len(row_numbers) > 1
        )
        return PreparedExcelImport(
            source_path=source_path,
            source_sha256=_file_sha256(source_path),
            records=tuple(records_by_gtin.values()),
            present_fields=present_fields,
            has_short_col="short_name" in columns,
            blank_rows=blank_rows,
            duplicates=duplicates,
            unknown_headers=tuple(unknown_headers),
            errors=tuple(errors),
        )
    finally:
        workbook.close()


def select_automatic_migration_source(
    base_dir: str | os.PathLike[str], database_path: str | os.PathLike[str]
) -> Path | None:
    if Path(database_path).exists():
        return None
    root = Path(base_dir)
    primary = root / PRIMARY_CATALOG_FILENAME
    if primary.is_file():
        return primary
    legacy = root / LEGACY_CATALOG_FILENAME
    return legacy if legacy.is_file() else None


def apply_prepared_import(
    store: CatalogStore,
    prepared: PreparedExcelImport,
    *,
    expected_revision: int | None = None,
    create_backup: bool = True,
) -> ImportResult:
    records = prepared.require_valid()
    return store.apply_import(
        records,
        present_fields=prepared.present_fields,
        has_short_col=prepared.has_short_col,
        expected_revision=expected_revision,
        create_backup=create_backup,
    )


def _record_value(record: CatalogProduct | Mapping, field: str) -> str:
    if isinstance(record, Mapping):
        value = record.get(field, "")
    else:
        value = getattr(record, field, "")
    return "" if value is None else str(value)


def export_catalog_xlsx(
    records: Iterable[CatalogProduct | Mapping], path: str | os.PathLike[str]
) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel catalog export") from exc

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp_path = destination.with_name(
        f".{destination.stem}.{uuid.uuid4().hex}.tmp.xlsx"
    )
    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = "Catalog"
        worksheet.append(CATALOG_HEADERS)
        for record in records:
            worksheet.append([_record_value(record, field) for field in CATALOG_FIELDS])
        for row in worksheet.iter_rows():
            for cell in row:
                cell.data_type = "s"
                cell.number_format = "@"
        workbook.save(temp_path)
        workbook.close()
        os.replace(temp_path, destination)
    except Exception:
        workbook.close()
        temp_path.unlink(missing_ok=True)
        raise
    return destination
