# -*- coding: utf-8 -*-

import os
import csv
import re
import datetime as dt

from utils import only_digits, short_gtin, add_years, add_months, add_days

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def load_kontur_raw(csv_path: str):
    rows = []
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        rdr = csv.reader(f, delimiter="\t", quotechar='"')
        for parts in rdr:
            if not parts or all((p or "").strip() == "" for p in parts):
                continue
            parts = [(p or "").strip() for p in parts]
            while len(parts) < 3:
                parts.append("")
            dm, gtin, name = parts[0], parts[1], parts[2]
            if dm.strip().upper() == "DM" and (name.strip().upper() in ("NAME", "")):
                continue
            if not dm.strip():
                continue
            rows.append({"DM": dm, "GTIN": gtin, "NAME": name})
    return rows


def _norm(s):
    return re.sub(r"\s+", "", (s or "").strip().lower())


def _parse_shelf_life(text: str):
    """{'days'| 'months' | 'years' | 'weeks' | 'raw'}"""
    t = (text or "").strip().lower()
    res = {'days': None, 'months': None, 'years': None, 'weeks': None, 'raw': text}
    if not t:
        return res
    m = re.search(r"(\d+)", t)
    if not m:
        return res
    n = int(m.group(1))
    if any(w in t for w in ["год", "лет", "года"]):
        res['years'] = n
        return res
    if any(w in t for w in ["мес", "месяц", "месяцев", "месяца"]):
        res['months'] = n
        return res
    if any(w in t for w in ["нед", "неделя", "недели", "недель"]):
        res['weeks'] = n
        return res
    if any(w in t for w in ["сут", "дн"]):
        res['days'] = n
        return res
    # По умолчанию считаем, что это месяцы, если не указано иное
    res['months'] = n
    return res


def read_product_map(xlsx_path: str):
    """Читает: GTIN, Упаковка, Срок годности, Шаблон партии, ShortName."""
    mapping = {}
    if not xlsx_path or not os.path.isfile(xlsx_path) or not load_workbook:
        return mapping
    try:
        wb = load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        hdrs = [str(c.value or "").strip() for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(hdrs)}

        idx_gtin = col.get("GTIN") or col.get("ГТИН") or col.get("gtin")
        idx_pack = col.get("Упаковка") or col.get("УПАКОВКА")
        idx_shelf = col.get("Срок годности") or col.get("СРОК ГОДНОСТИ")
        idx_part = col.get("Шаблон партии") or col.get("ШАБЛОН ПАРТИИ") or col.get("Партия") or col.get("ПАРТИЯ")
        idx_short = (col.get("ShortName") or col.get("SHORTNAME") or col.get("Shortname") or col.get("shortname")
                     or col.get("Short Name") or col.get("SHORT NAME") or col.get("Короткое имя") or col.get("КОРОТКОЕ ИМЯ"))
        idx_brand = col.get("Бренд") or col.get("бренд") or col.get("Brand") or col.get("BRAND")
        idx_client = col.get("Клиент") or col.get("клиент") or col.get("Client") or col.get("CLIENT")
        idx_packaging = col.get("Упаковка отображаемая") or col.get("Упаковка для стикера") or col.get("PACKAGING") or col.get("Packaging")
        idx_weight = (col.get("Вес") or col.get("вес") or col.get("Weight") or col.get("WEIGHT")
                      or col.get("Объем") or col.get("объем") or col.get("Объём") or col.get("объём"))

        # Фолбэки по подстрокам в заголовке (ловим, например, "Вес, г" / "Weight, g")
        if not idx_weight:
            for i, h in enumerate(hdrs):
                hlow = h.lower()
                if ("вес" in hlow) or ("weight" in hlow) or ("объем" in hlow) or ("объём" in hlow):
                    idx_weight = i + 1
                    break
        if not idx_brand:
            for i, h in enumerate(hdrs):
                hlow = h.lower()
                if "brand" in hlow or "бренд" in hlow:
                    idx_brand = i + 1
                    break
        if not idx_packaging:
            for i, h in enumerate(hdrs):
                hlow = h.lower()
                if "pack" in hlow or "упаков" in hlow:
                    idx_packaging = i + 1
                    break

        def cell(row, j):
            if not j: return ""
            v = row[j - 1].value
            return "" if v is None else str(v).strip()

        has_short = bool(idx_short)
        for row in ws.iter_rows(min_row=2):
            g = only_digits(cell(row, idx_gtin))
            if not g:
                continue
            pack = _norm(cell(row, idx_pack))
            shelf = cell(row, idx_shelf)
            fmt = ""
            if "ведро" in pack:
                fmt = "30x20"
            elif any(w in pack for w in ("банка", "туба")):
                fmt = "16x16"
            shelf_parsed = _parse_shelf_life(shelf)
            part_tpl = cell(row, idx_part)
            short_from_xlsx = cell(row, idx_short)
            mapping[g] = {
                "FORMAT": fmt,
                "SHELF": shelf_parsed,
                "PART_TEMPLATE": part_tpl,
                "SHORTNAME": short_from_xlsx,
                "BRAND": cell(row, idx_brand),
                "CLIENT": cell(row, idx_client),
                "PACKAGING": cell(row, idx_packaging) or cell(row, idx_pack),
                "WEIGHT": cell(row, idx_weight),
            }
        mapping["_HAS_SHORT_COL"] = has_short
        return mapping
    except Exception:
        return {}


def choose_format_for(gtin, product_map, manual_choice):
    if manual_choice in ("16x16", "30x20"):
        return manual_choice
    info = product_map.get(only_digits(gtin), {})
    fmt = (info.get("FORMAT") or "") if info else ""
    return fmt if fmt in ("16x16", "30x20") else "16x16"


def make_part_num(prod_date, part_template: str | None):
    yymmdd = prod_date.strftime("%y%m%d") if prod_date else ""
    if part_template:
        return part_template.replace("{DATE}", yymmdd)
    return yymmdd


def enrich_row(base_row, idx1, prod_date, exp_days_override, product_map, mode_choice):
    dm, gtin, name = base_row.get("DM", ""), base_row.get("GTIN", ""), base_row.get("NAME", "")
    fmt = choose_format_for(gtin, product_map, mode_choice)
    info = product_map.get(only_digits(gtin), {})
    pd = prod_date or dt.date.today()

    exp_date = None
    shelf_info = info.get("SHELF") or {}
    log_shelf = ""
    if shelf_info:
        if shelf_info.get('years'): exp_date = add_years(pd, int(shelf_info['years'])); log_shelf = f"years={shelf_info['years']}"
        elif shelf_info.get('months'): exp_date = add_months(pd, int(shelf_info['months'])); log_shelf = f"months={shelf_info['months']}"
        elif shelf_info.get('weeks'): exp_date = add_days(pd, int(shelf_info['weeks']) * 7); log_shelf = f"weeks={shelf_info['weeks']}"
        elif shelf_info.get('days'): exp_date = add_days(pd, int(shelf_info['days'])); log_shelf = f"days={shelf_info['days']}"
    if (exp_date is None) and (exp_days_override is not None):
        exp_date = add_days(pd, int(exp_days_override)); log_shelf = f"override_days={exp_days_override}"

    part_num = make_part_num(pd, info.get("PART_TEMPLATE") or "")
    short_excel = (info.get("SHORTNAME") or "").strip()

    return {
        "DM": dm, "GTIN": gtin, "NAME": name,
        "BRAND": info.get("BRAND", ""),
        "CLIENT": info.get("CLIENT", ""),
        "PACKAGING": info.get("PACKAGING", ""),
        "WEIGHT": info.get("WEIGHT", ""),
        "ShortGTIN": short_gtin(gtin),
        "ShortName": short_excel if short_excel else name,
        "PROD_DATE": pd.strftime("%d.%m.%Y") if pd else "",
        "NUM": str(idx1),
        "EXP_DATE": exp_date.strftime("%d.%m.%Y") if exp_date else "",
        "PART_NUM": part_num,
        "_FORMAT": fmt,
        "_SHELF_LOG": log_shelf,
        "_SHORT_SRC": "Excel" if short_excel else "NAME"
    }
